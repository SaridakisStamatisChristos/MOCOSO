"""Compile all manuscript formats from manuscript.md and references.bib.

Requires pandoc with citeproc, xelatex, and python-docx.
Run from any directory. This script never changes research solver files.
"""
from pathlib import Path
import os
import re
import subprocess

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "exports"
OUT.mkdir(exist_ok=True)
source = (ROOT / "manuscript.md").read_text()
bib = (ROOT / "references.bib").read_text()
cited = set(re.findall(r"@([a-z]+[0-9]+)", source))
entries = set(re.findall(r"@\w+\{([^,]+),", bib))
assert cited == entries, ("Citation mismatch", cited, entries)

common = ["pandoc", "manuscript.md", "--from=markdown", "--standalone",
          "--citeproc", "--bibliography=references.bib"]

def run(args):
    subprocess.run(args, cwd=ROOT, check=True)

# Pandoc creates native OMML equations; python-docx adjusts styles without
# replacing the math objects. The LaTeX and PDF use the same parsed manuscript.
run(common + ["-o", "exports/manuscript.docx"])
doc = Document(OUT / "manuscript.docx")
for section in doc.sections:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.85)
    section.left_margin = section.right_margin = Inches(1)
for style in doc.styles:
    if style.type == 1:
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
for name in ("Normal", "Body Text", "First Paragraph"):
    if name in doc.styles:
        st = doc.styles[name]
        st.font.size = Pt(11)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.05
styles_by_name = {s.name: s for s in doc.styles}
for name, size in (("Title", 19), ("Heading 1", 14), ("Heading 2", 12)):
    style = styles_by_name[name]
    style.font.size = Pt(size)
    style.font.bold = True
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(6)
for table in doc.tables:
    table.autofit = False
    widths = ([2.7, 0.8, 3.0] if len(table.rows) == 7 else [2.1, 2.2, 2.2])
    for r, row in enumerate(table.rows):
        if r == 0:
            repeat = OxmlElement("w:tblHeader")
            row._tr.get_or_add_trPr().append(repeat)
        for c, cell in enumerate(row.cells):
            cell.width = Inches(widths[c])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                border = OxmlElement("w:" + side)
                for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    border.set(qn("w:"+key), value)
                borders.append(border)
            props.append(borders)
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                margin = OxmlElement("w:" + side)
                margin.set(qn("w:w"), "80")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            props.append(margins)
            if r == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E8EDF1")
                props.append(shade)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.space_before = Pt(3)
                for run_ in paragraph.runs:
                    run_.font.size = Pt(10)
                    if r == 0:
                        run_.font.bold = True
for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
doc.core_properties.title = "Canonical witnesses for modular convex subsequences with a small index deficit"
doc.core_properties.author = "Stamatis-Christos Saridakis"
doc.core_properties.subject = "Research draft; corrected G9 theorem and audit"
doc.save(OUT / "manuscript.docx")

run(common + ["--to=latex", "-o", "exports/manuscript.tex"])
run(common + ["--pdf-engine=xelatex", "-o", "exports/manuscript.pdf"])
print("Built DOCX with native math, standalone LaTeX, and PDF from one source.")
